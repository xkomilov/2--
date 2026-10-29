#!/usr/bin/env python3
"""
Build per-topic .xlsx test tables in the template format.

Template format (shablon.xlsx):
- Row 1 header: Question | A | B | C | D | E | Correct | Comment
- Text wrap + alignment on data cells
- Data validation on Correct column (list: A,B,C,D,E)
- NO topic name written into the sheet (task requirement)

Input: a JSON file (questions.json) with structure:
{
  "001": [
     {"q": "...", "A": "...", "B": "...", ... "E": "...",
      "correct": "C", "comment": "To'g'ri javob: ...  Manba: ..."},
     ...
  ],
  ...
}

Usage: python3 build_xlsx.py questions.json out_dir
"""
import sys, json, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), 'pylibs'))
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

HEADER = ["Question", "A", "B", "C", "D", "E", "Correct", "Comment"]

# Column widths tuned for readability
WIDTHS = {"A": 46, "B": 26, "C": 26, "D": 26, "E": 26, "F": 26, "G": 9, "H": 60}

def build_workbook(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test"

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Header
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    header_font = Font(bold=True)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for ci, name in enumerate(HEADER, start=1):
        c = ws.cell(row=1, column=ci, value=name)
        c.fill = header_fill
        c.font = header_font
        c.alignment = header_align
        c.border = border

    # Data rows
    data_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for ri, row in enumerate(rows, start=2):
        values = [row["q"], row["A"], row["B"], row["C"], row["D"], row["E"],
                  row["correct"], row["comment"]]
        for ci, val in enumerate(values, start=1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.border = border
            if ci == 7:  # Correct column
                c.alignment = center_align
            else:
                c.alignment = data_align

    # Data validation on Correct column (G), rows 2..N
    n = len(rows)
    dv = DataValidation(type="list", formula1='"A,B,C,D,E"', allow_blank=False,
                        showDropDown=False)
    dv.error = "Faqat A, B, C, D yoki E bo'lishi mumkin"
    dv.errorTitle = "Noto'g'ri qiymat"
    dv.prompt = "To'g'ri javob ustunini tanlang (A-E)"
    dv.promptTitle = "Correct"
    ws.add_data_validation(dv)
    if n >= 1:
        dv.add(f"G2:G{n+1}")

    # Column widths
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w

    # Freeze header
    ws.freeze_panes = "A2"
    return wb

def main():
    qfile = sys.argv[1]
    outdir = sys.argv[2]
    os.makedirs(outdir, exist_ok=True)
    with open(qfile, encoding="utf-8") as f:
        data = json.load(f)
    for topic, rows in sorted(data.items()):
        wb = build_workbook(rows)
        path = os.path.join(outdir, f"{topic}.xlsx")
        wb.save(path)
        print(f"wrote {path} ({len(rows)} questions)")

if __name__ == "__main__":
    main()
